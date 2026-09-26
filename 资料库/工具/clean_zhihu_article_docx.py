#!/usr/bin/env python3
"""Clean the user-supplied Zhihu article captures for source URLs 03 and 06."""

from dataclasses import dataclass
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]


@dataclass(frozen=True)
class Article:
    source: Path
    clean: Path
    markdown: Path
    source_url: str
    title_paragraph: int
    body_ranges: tuple[tuple[int, int], ...]


ARTICLES = (
    Article(
        source=ROOT / "02_网页原始内容/03_transformer_用户抓取/03_transformer_用户抓取原稿.docx",
        clean=ROOT / "02_网页原始内容/03_transformer_用户抓取/03_transformer_知乎文章_清洗版.docx",
        markdown=ROOT / "03_可检索正文/03_transformer_知乎文章_用户抓取清洗正文.md",
        source_url="https://zhuanlan.zhihu.com/p/53682800",
        title_paragraph=29,
        # Intro, then the article body through its reference list.
        body_ranges=((38, 38), (41, 163)),
    ),
    Article(
        source=ROOT / "02_网页原始内容/06_diffusion_tutorial_用户抓取/06_diffusion_用户抓取原稿.docx",
        clean=ROOT / "02_网页原始内容/06_diffusion_tutorial_用户抓取/06_diffusion_知乎文章_清洗版.docx",
        markdown=ROOT / "03_可检索正文/06_diffusion_知乎文章_用户抓取清洗正文.md",
        source_url="https://zhuanlan.zhihu.com/p/525106459",
        title_paragraph=29,
        # Article begins at 前言 and ends at the reference list.
        body_ranges=((48, 156),),
    ),
)


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def has_numbering(paragraph: Paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def image_count(paragraph: Paragraph) -> int:
    return len(paragraph._p.xpath(".//w:drawing")) + len(paragraph._p.xpath(".//w:pict"))


def insert_source_note_after(title: Paragraph, source_url: str) -> None:
    note = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = f"来源：{source_url}（用户抓取稿清洗；网页导航、评论和推荐内容已移除）"
    run.append(text)
    note.append(run)
    title._p.addnext(note)


def remove_broken_images(document: Document) -> int:
    removed = 0
    valid_relationships = set(document.part.rels)
    for blip in list(document.element.body.xpath(".//a:blip")):
        relationship_id = blip.get(qn("r:embed"))
        if not relationship_id or relationship_id in valid_relationships:
            continue
        drawing = blip
        while drawing is not None and drawing.tag not in {qn("w:drawing"), qn("w:pict")}:
            drawing = drawing.getparent()
        if drawing is not None and drawing.getparent() is not None:
            drawing.getparent().remove(drawing)
            removed += 1
    return removed


def remove_invalid_style_nodes(document: Document) -> int:
    removed = 0
    for node in list(document.styles.element.xpath(".//w:uiPriority")):
        node.getparent().remove(node)
        removed += 1
    return removed


def keep_indices(article: Article) -> set[int]:
    keep = {article.title_paragraph}
    for start, end in article.body_ranges:
        keep.update(range(start, end + 1))
    return keep


def make_clean_docx(article: Article) -> tuple[int, int]:
    copy2(article.source, article.clean)
    document = Document(article.clean)
    original = list(document.paragraphs)
    keep = keep_indices(article)
    for index, paragraph in enumerate(original, 1):
        if index not in keep:
            remove_paragraph(paragraph)
    insert_source_note_after(document.paragraphs[0], article.source_url)
    broken_images = remove_broken_images(document)
    style_nodes = remove_invalid_style_nodes(document)
    document.save(article.clean)
    return broken_images, style_nodes


def make_markdown(article: Article) -> tuple[int, int, int]:
    document = Document(article.clean)
    lines = [
        "<!-- 由用户提供的网页抓取稿清洗生成；正文未作知识性改写。 -->",
        "",
    ]
    figure_number = 0
    in_code = False

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        is_code = style == "HTML Preformatted"

        if is_code and not in_code:
            lines.append("```python")
            in_code = True
        elif not is_code and in_code:
            lines.extend(["```", ""])
            in_code = False

        if is_code:
            lines.append(paragraph.text.rstrip())
        elif text:
            if style.lower().startswith("heading"):
                try:
                    level = int(style.split()[-1])
                except ValueError:
                    level = 2
                lines.extend([f"{'#' * level} {text}", ""])
            elif has_numbering(paragraph):
                lines.extend([f"- {text}", ""])
            else:
                lines.extend([text, ""])

        for _ in range(image_count(paragraph)):
            figure_number += 1
            lines.extend([f"[插图 {figure_number}：保存在清洗版 Word 文档中]", ""])

    if in_code:
        lines.extend(["```", ""])

    article.markdown.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    characters = sum(len(p.text) for p in document.paragraphs)
    return figure_number, characters, len(document.paragraphs)


if __name__ == "__main__":
    for article in ARTICLES:
        broken, styles = make_clean_docx(article)
        figures, characters, paragraphs = make_markdown(article)
        print(article.clean.name)
        print(f"  removed_broken_images={broken} removed_invalid_style_nodes={styles}")
        print(f"  paragraphs={paragraphs} figures={figures} characters={characters}")
        print(f"  markdown={article.markdown}")
